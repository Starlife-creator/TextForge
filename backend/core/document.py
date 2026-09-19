from docx import Document
from pathlib import Path
import tempfile, os

def write_text_docx(text: str, output_path: str):
    text = text.replace('\r\n', '\n').replace('\r', '\n')
    doc = Document()
    for para in text.split("\n\n"):
        p = doc.add_paragraph()
        for line in para.split("\n"):
            p.add_run(line)
            p.add_run().add_break()
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    fd, tmp = tempfile.mkstemp(dir=output_path.parent, prefix=output_path.stem + "_", suffix=".docx")
    os.close(fd)
    try:
        doc.save(tmp)
        os.replace(tmp, output_path)
    except Exception:
        if os.path.exists(tmp): os.unlink(tmp)
        raise
