"""docx 输出：普通模式保持旧行为，富文本模式升级标题/分页/粗斜体。"""

import tempfile
import zipfile
from pathlib import Path

from docx import Document

from core.document import write_chapters_docx


def _xml(path: Path) -> str:
    with zipfile.ZipFile(path) as z:
        return z.read("word/document.xml").decode("utf-8")


def test_rich_docx_adds_headings_page_break_and_inline_marks():
    with tempfile.TemporaryDirectory() as tmp:
        path = Path(tmp) / "book.docx"
        write_chapters_docx(
            [
                "第一章 开端\n\n这是**加粗**与*斜体*的正文。",
                "第二章 继续\n\n普通段落。",
            ],
            str(path),
            rich=True,
            title="测试书",
        )
        doc = Document(str(path))
        styles = [p.style.name for p in doc.paragraphs]
        assert styles.count("Heading 1") == 2
        assert any(run.bold for p in doc.paragraphs for run in p.runs)
        assert any(run.italic for p in doc.paragraphs for run in p.runs)
        assert 'w:type="page"' in _xml(path)
        assert doc.core_properties.title == "测试书"


def test_rich_docx_renders_markdown_heading():
    with tempfile.TemporaryDirectory() as tmp:
        path = Path(tmp) / "md.docx"
        write_chapters_docx(["## 小节标题\n\n正文。"], str(path), rich=True)
        doc = Document(str(path))
        assert any(p.style.name == "Heading 2" for p in doc.paragraphs)


def test_rich_docx_does_not_style_long_paragraph_as_heading():
    with tempfile.TemporaryDirectory() as tmp:
        path = Path(tmp) / "long.docx"
        line = "第一章 " + "这是一段很长的正文。" * 10
        write_chapters_docx([line + "\n\n后文。"], str(path), rich=True)
        doc = Document(str(path))
        assert all(p.style.name != "Heading 1" for p in doc.paragraphs)


def test_plain_docx_keeps_legacy_layout():
    with tempfile.TemporaryDirectory() as tmp:
        path = Path(tmp) / "plain.docx"
        write_chapters_docx(
            ["第一章 开端\n\n正文一。", "第二章 继续\n\n正文二。"],
            str(path),
        )
        doc = Document(str(path))
        assert all(p.style.name == "Normal" for p in doc.paragraphs)
        assert "---" in [p.text for p in doc.paragraphs]
        assert 'w:type="page"' not in _xml(path)
